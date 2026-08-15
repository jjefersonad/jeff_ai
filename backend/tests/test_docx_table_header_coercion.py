"""Tolerância: header como lista de colunas em create_docx_document."""
from __future__ import annotations

from pathlib import Path
from unittest.mock import AsyncMock

import src.tools.create_docx_document_tool as docx_tool
from docx import Document as DocxReader
from src.models.docx_document import DocxBlockInput, DocxDocumentInput


def test_docx_block_coerces_header_list_into_rows() -> None:
    block = DocxBlockInput(
        type="table",
        header=["Mês", "Receita"],
        rows=[["Jan", "12000"]],
    )
    assert block.header is True
    assert block.rows == [["Mês", "Receita"], ["Jan", "12000"]]


def test_docx_block_does_not_duplicate_header_when_already_in_rows() -> None:
    block = DocxBlockInput(
        type="table",
        header=["Mês", "Receita"],
        rows=[["Mês", "Receita"], ["Jan", "12000"]],
    )
    assert block.header is True
    assert block.rows == [["Mês", "Receita"], ["Jan", "12000"]]


async def test_tool_accepts_header_as_column_name_list(monkeypatch, tmp_path):
    """Reproduz o payload típico do LLM que antes falhava na validação."""
    monkeypatch.setattr(docx_tool, "require_user_docs_dir", AsyncMock(return_value=tmp_path))
    monkeypatch.setattr(docx_tool, "_document_url_prefix", lambda: "/api/files")
    monkeypatch.setattr(docx_tool, "record_ownership", AsyncMock())

    result = await docx_tool.create_docx_document.coroutine(
        DocxDocumentInput(
            title="Agendamentos",
            blocks=[
                DocxBlockInput(type="heading", text="Lista", level=1),
                DocxBlockInput(
                    type="table",
                    header=["ID", "Data", "Status", "Mensagem"],
                    rows=[
                        ["abc", "07/08/2026", "Sucesso", "ok"],
                    ],
                ),
            ],
        )
    )

    assert "error" not in result
    assert Path(result["path"]).is_file()
    doc = DocxReader(result["path"])
    assert len(doc.tables) == 1
    table = doc.tables[0]
    assert table.cell(0, 0).text == "ID"
    assert table.cell(0, 3).text == "Mensagem"
    assert table.cell(1, 2).text == "Sucesso"
